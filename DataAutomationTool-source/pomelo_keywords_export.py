"""Download Pomelo keyword API responses in batches.

This CLI replays a locally saved DevTools "Copy as cURL" request,
replaces its keyword payload in batches, and saves each untouched response as
one JSONL record. Sensitive headers and cookies stay local and are never
printed or copied into the output.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shlex
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable

try:
    import openpyxl
except Exception:  # pragma: no cover - checked at runtime
    openpyxl = None

from output_writer import ROOT, beijing_timestamp

DEFAULT_OUTPUT_ROOT = ROOT / "output" / "Step9_Pomelo_Keywords"
RAW_JSONL_NAME = "pomelo_keywords_raw_pages.jsonl"
RUN_META_NAME = "pomelo_keywords_run_meta.json"
DEFAULT_KEYWORDS_SHEET = "关键词"


@dataclass(frozen=True)
class CurlRequest:
    url: str
    headers: dict[str, str]
    payload: dict[str, Any]


class BatchIncompleteError(RuntimeError):
    def __init__(self, failed_batch: int, completed_batch: int, total_batches: int, raw_path: Path) -> None:
        super().__init__(f"step9_batch_incomplete:{failed_batch}")
        self.failed_batch = failed_batch
        self.completed_batch = completed_batch
        self.total_batches = total_batches
        self.raw_path = raw_path


def require_openpyxl() -> None:
    if openpyxl is None:
        raise RuntimeError("缺少 openpyxl，无法读取 Excel 关键词文件。请先安装依赖。")


def normalize_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def normalize_key(value: Any) -> str:
    return normalize_text(value).lower()


def decode_windows_curl(text: str) -> list[str]:
    normalized = text.replace("^\r\n", " ").replace("^\n", " ")
    normalized = normalized.replace("^", "")
    return shlex.split(normalized, posix=True)


def load_curl(path: str | Path) -> CurlRequest:
    curl_path = Path(path)
    if not curl_path.exists():
        raise RuntimeError(f"cURL 文件不存在：{curl_path}")
    parts = decode_windows_curl(curl_path.read_text(encoding="utf-8"))
    if not parts or parts[0].lower() != "curl":
        raise RuntimeError("cURL 文件格式不正确：应以 curl 开头。")

    url = ""
    headers: dict[str, str] = {}
    data_raw = ""
    index = 1
    while index < len(parts):
        token = parts[index]
        if token in ("-H", "--header") and index + 1 < len(parts):
            header = parts[index + 1]
            if ":" in header:
                name, value = header.split(":", 1)
                headers[name.strip()] = value.strip()
            index += 2
            continue
        if token in ("-b", "--cookie") and index + 1 < len(parts):
            headers["Cookie"] = parts[index + 1]
            index += 2
            continue
        if token == "--data-raw" and index + 1 < len(parts):
            data_raw = parts[index + 1]
            index += 2
            continue
        if not token.startswith("-") and not url:
            url = token
        index += 1

    if not url:
        raise RuntimeError("cURL 中未找到请求 URL。")
    if not data_raw:
        raise RuntimeError("cURL 中未找到 --data-raw 请求体。")
    try:
        payload = json.loads(data_raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"无法解析 cURL JSON 请求体：{exc}") from exc
    return CurlRequest(url=url, headers=headers, payload=payload)


def find_column(headers: list[str], candidates: Iterable[str]) -> int:
    normalized_headers = [normalize_key(header) for header in headers]
    for candidate in candidates:
        normalized_candidate = normalize_key(candidate)
        for index, header in enumerate(normalized_headers):
            if header == normalized_candidate:
                return index
    for candidate in candidates:
        normalized_candidate = normalize_key(candidate)
        for index, header in enumerate(normalized_headers):
            if normalized_candidate and normalized_candidate in header:
                return index
    raise RuntimeError(f"找不到列：{', '.join(candidates)}")


def dedupe_keywords(values: Iterable[Any]) -> list[str]:
    keywords: list[str] = []
    seen: set[str] = set()
    for value in values:
        keyword = normalize_text(value)
        key = keyword.lower()
        if not keyword or key in seen:
            continue
        seen.add(key)
        keywords.append(keyword)
    return keywords


def keywords_from_payload(payload: dict[str, Any]) -> list[str]:
    return dedupe_keywords(payload.get("resource", {}).get("searchTerms", []))


def keywords_from_text(path: Path) -> list[str]:
    return dedupe_keywords(path.read_text(encoding="utf-8").splitlines())


def keywords_from_csv(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        rows = list(reader)
    if not rows:
        return []
    headers = [normalize_text(value) for value in rows[0]]
    try:
        keyword_index = find_column(headers, ["关键词", "keyword", "keywords", "searchTerm"])
        return dedupe_keywords(row[keyword_index] if keyword_index < len(row) else "" for row in rows[1:])
    except RuntimeError:
        return dedupe_keywords(row[0] if row else "" for row in rows)


def keywords_from_excel(path: Path, sheet_name: str) -> list[str]:
    require_openpyxl()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return []
        headers = [normalize_text(value) for value in header_row]
        keyword_index = find_column(headers, ["关键词", "keyword", "keywords", "searchTerm"])
        return dedupe_keywords(row[keyword_index] if keyword_index < len(row) else "" for row in rows_iter)
    finally:
        workbook.close()


def load_keywords(keywords_path: str | Path | None, curl_payload: dict[str, Any], sheet_name: str, limit: int | None) -> list[str]:
    if keywords_path:
        path = Path(keywords_path)
        if not path.exists():
            raise RuntimeError(f"关键词文件不存在：{path}")
        suffix = path.suffix.lower()
        if suffix in (".xlsx", ".xlsm"):
            keywords = keywords_from_excel(path, sheet_name)
        elif suffix == ".csv":
            keywords = keywords_from_csv(path)
        else:
            keywords = keywords_from_text(path)
    else:
        keywords = keywords_from_payload(curl_payload)
    if limit is not None:
        if limit <= 0:
            raise RuntimeError("--limit 必须大于 0。")
        keywords = keywords[:limit]
    if not keywords:
        raise RuntimeError("没有可处理的关键词。")
    return keywords


def make_batches(values: list[str], batch_size: int, max_batches: int | None) -> list[list[str]]:
    if batch_size <= 0:
        raise RuntimeError("--batch-size 必须大于 0。")
    batches = [values[index : index + batch_size] for index in range(0, len(values), batch_size)]
    if max_batches is not None:
        if max_batches <= 0:
            raise RuntimeError("--max-batches 必须大于 0。")
        batches = batches[:max_batches]
    return batches


def keyword_fingerprint(keywords: list[str]) -> str:
    digest = hashlib.sha256()
    for keyword in keywords:
        digest.update(normalize_text(keyword).encode("utf-8"))
        digest.update(b"\0")
    return digest.hexdigest()


def build_run_meta(keywords: list[str], batch_size: int, country: str, asins_count: int) -> dict[str, Any]:
    return {
        "country": country,
        "batch_size": batch_size,
        "asins_count": asins_count,
        "keyword_count": len(keywords),
        "keyword_fingerprint": keyword_fingerprint(keywords),
    }


def write_run_meta(path: Path, meta: dict[str, Any]) -> None:
    path.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def load_run_meta(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise RuntimeError(f"续跑需要元信息文件，但未找到：{path}")
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"无法解析续跑元信息文件：{path}") from exc
    if not isinstance(data, dict):
        raise RuntimeError(f"续跑元信息文件格式不正确：{path}")
    return data


def validate_resume_meta(existing: dict[str, Any], current: dict[str, Any]) -> None:
    fields = ["country", "batch_size", "asins_count", "keyword_count", "keyword_fingerprint"]
    mismatches = [field for field in fields if existing.get(field) != current.get(field)]
    if mismatches:
        raise RuntimeError(
            "续跑参数不一致："
            + ", ".join(mismatches)
            + "。请使用原始关键词文件和参数，或换一个新的 output-dir。"
        )


def iter_raw_records(raw_path: Path) -> Iterable[dict[str, Any]]:
    if not raw_path.exists():
        return
    with raw_path.open("r", encoding="utf-8") as file:
        for line_no, line in enumerate(file, start=1):
            text = line.strip()
            if not text:
                continue
            try:
                record = json.loads(text)
            except json.JSONDecodeError:
                print(f"警告：raw JSONL 第 {line_no} 行无法解析，已跳过。")
                continue
            if not isinstance(record, dict):
                print(f"警告：raw JSONL 第 {line_no} 行不是对象，已跳过。")
                continue
            yield record


def completed_batches_from_raw(raw_path: Path) -> set[int]:
    completed: set[int] = set()
    for record in iter_raw_records(raw_path):
        batch_no = record.get("batch")
        response = record.get("response")
        if isinstance(batch_no, int) and isinstance(response, dict):
            completed.add(batch_no)
    return completed


def payload_for_batch(template: dict[str, Any], country: str, keywords: list[str], asins_count: int) -> dict[str, Any]:
    payload = json.loads(json.dumps(template, ensure_ascii=False))
    payload.setdefault("resource", {})["country"] = country
    payload.setdefault("resource", {})["searchTerms"] = keywords
    payload.setdefault("biz", {})["entities"] = [{"country": country, "searchTerm": keyword} for keyword in keywords]
    payload.setdefault("biz", {})["asinsCount"] = asins_count
    return payload


def request_headers(headers: dict[str, str]) -> dict[str, str]:
    result = dict(headers)
    if "Content-Type" not in result and "content-type" not in result:
        result["Content-Type"] = "application/json"
    return result


def post_json(request: CurlRequest, payload: dict[str, Any], timeout: int) -> dict[str, Any]:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(request.url, data=body, headers=request_headers(request.headers), method="POST")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as response:
            text = response.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        raise RuntimeError(f"HTTP {exc.code}：柚子关键词请求失败。") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"网络请求失败：{exc.reason}") from exc
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"响应不是合法 JSON：{exc}") from exc


def validate_response(data: dict[str, Any]) -> None:
    if not isinstance(data, dict):
        raise RuntimeError("柚子关键词响应不是 JSON 对象。")
    entities = data.get("entities")
    if entities is not None and not isinstance(entities, list):
        raise RuntimeError("柚子关键词响应结构异常。")


def validate_run_size(total_keywords: int, max_batches: int | None, confirm_full: bool) -> None:
    if max_batches is None and total_keywords > 200 and not confirm_full:
        raise RuntimeError(
            f"将处理 {total_keywords} 个关键词。为避免误跑全量，请传 --max-batches 做测试，或显式传 --confirm-full。"
        )


def run_export(
    curl_path: Path,
    keywords_path: str | None,
    sheet_name: str,
    batch_size: int,
    max_batches: int | None,
    limit: int | None,
    delay: float,
    output_dir: Path,
    country: str,
    asins_count: int,
    timeout: int,
    confirm_full: bool,
    resume: bool = False,
    progress_callback: Callable[[str], None] | None = None,
    save_checkpoint: Callable[[dict[str, object]], None] | None = None,
    stop_requested: Callable[[], bool] | None = None,
    publish_progress: Callable[[int, int], None] | None = None,
) -> Path:
    def emit_progress(message: str) -> None:
        if progress_callback:
            progress_callback(message)
        else:
            print(message)

    request = load_curl(curl_path)
    keywords = load_keywords(keywords_path, request.payload, sheet_name, limit)
    validate_run_size(len(keywords), max_batches, confirm_full)
    batches = make_batches(keywords, batch_size, max_batches)

    output_dir.mkdir(parents=True, exist_ok=True)
    raw_path = output_dir / RAW_JSONL_NAME
    meta_path = output_dir / RUN_META_NAME
    current_meta = build_run_meta(keywords, batch_size, country, asins_count)
    if resume:
        existing_meta = load_run_meta(meta_path)
        validate_resume_meta(existing_meta, current_meta)
        completed_batches = completed_batches_from_raw(raw_path)
    else:
        write_run_meta(meta_path, current_meta)
        completed_batches = set()

    if save_checkpoint:
        save_checkpoint({
            "completedBatch": len(completed_batches),
            "totalBatches": len(batches),
            "outputDir": str(output_dir),
        })
    if publish_progress:
        publish_progress(len(completed_batches), len(batches))
    pending_batches = [(batch_no, batch_keywords) for batch_no, batch_keywords in enumerate(batches, start=1) if batch_no not in completed_batches]
    emit_progress(
        f"准备请求：关键词 {len(keywords)} 个；实际批次 {len(batches)}；"
        f"已完成 {len(completed_batches)}；待请求 {len(pending_batches)}；"
        f"batch-size={batch_size}；asinsCount={asins_count}"
    )
    raw_mode = "a" if resume else "w"
    incomplete_batch: int | None = None
    with raw_path.open(raw_mode, encoding="utf-8") as raw_file:
        for batch_no, batch_keywords in pending_batches:
            if stop_requested and stop_requested():
                break
            try:
                payload = payload_for_batch(request.payload, country, batch_keywords, asins_count)
                data = post_json(request, payload, timeout=timeout)
                validate_response(data)
                raw_file.write(json.dumps({"batch": batch_no, "keywords": batch_keywords, "response": data}, ensure_ascii=False) + "\n")
                raw_file.flush()
                os.fsync(raw_file.fileno())
                checkpoint = {
                    "completedBatch": batch_no,
                    "totalBatches": len(batches),
                    "outputDir": str(output_dir),
                }
                if save_checkpoint:
                    save_checkpoint(checkpoint)
                if publish_progress:
                    publish_progress(batch_no, len(batches))
                emit_progress(f"批次 {batch_no}/{len(batches)} 成功：已保存原始响应")
            except Exception as exc:
                message = str(exc)
                emit_progress(f"批次 {batch_no}/{len(batches)} 失败：{message}")
                incomplete_batch = batch_no
                break
            if delay > 0 and pending_batches and batch_no != pending_batches[-1][0]:
                time.sleep(delay)

    write_run_meta(meta_path, current_meta)
    emit_progress(f"Raw JSONL 输出：{raw_path}")
    if incomplete_batch is not None:
        completed_batch = max(completed_batches_from_raw(raw_path), default=0)
        raise BatchIncompleteError(incomplete_batch, completed_batch, len(batches), raw_path)
    return raw_path


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="下载柚子关键词原始 JSONL")
    parser.add_argument("--curl", required=True, help="DevTools Copy as cURL 保存的请求档案")
    parser.add_argument("--keywords", help="关键词文件：xlsx、csv 或 txt；不传则使用 cURL payload.resource.searchTerms")
    parser.add_argument("--sheet", default=DEFAULT_KEYWORDS_SHEET, help=f"Excel sheet 名，默认 {DEFAULT_KEYWORDS_SHEET}")
    parser.add_argument("--batch-size", type=int, default=100, help="每批关键词数，默认 100")
    parser.add_argument("--max-batches", type=int, help="最多请求批次数；测试建议传 1 或 2")
    parser.add_argument("--limit", type=int, help="最多读取多少个关键词")
    parser.add_argument("--delay", type=float, default=2.0, help="批次间隔秒数")
    parser.add_argument("--country", default="US", help="国家，默认 US")
    parser.add_argument("--asins-count", type=int, default=50, help="每个关键词返回的 ASIN 数，默认 50")
    parser.add_argument("--timeout", type=int, default=120, help="单批请求超时秒数")
    parser.add_argument("--output-dir", help="输出目录，默认 output/Step9_Pomelo_Keywords/时间戳")
    parser.add_argument("--confirm-full", action="store_true", help="确认执行超过 200 个关键词的全量任务")
    parser.add_argument("--resume", action="store_true", help="从已有 output-dir 续跑，跳过 raw JSONL 中已成功的批次")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if args.resume and not args.output_dir:
        print("运行失败：--resume 必须配合明确的 --output-dir 使用。")
        return 1
    try:
        output_dir = Path(args.output_dir) if args.output_dir else DEFAULT_OUTPUT_ROOT / beijing_timestamp()
        run_export(
            curl_path=Path(args.curl),
            keywords_path=args.keywords,
            sheet_name=args.sheet,
            batch_size=args.batch_size,
            max_batches=args.max_batches,
            limit=args.limit,
            delay=args.delay,
            output_dir=output_dir,
            country=args.country.upper(),
            asins_count=args.asins_count,
            timeout=args.timeout,
            confirm_full=args.confirm_full,
            resume=args.resume,
        )
        return 0
    except Exception as exc:
        print(f"运行失败：{exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
