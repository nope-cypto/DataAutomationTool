"""Wheat ASIN traffic-extension raw response downloader.

Minimal CLI:

    Electron/Worker invokes this module directly for the supported product flow.

The fetch command uses a cURL copied from the browser Network panel. It does not
click the web page; it replays the same authorized JSON request page by page.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import os
import shlex
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

try:
    import openpyxl
except Exception:  # pragma: no cover - checked at runtime
    openpyxl = None

from output_writer import create_run_output_dir
from web_workbench.worker.security import credential_free_curl_fingerprint

RAW_PAGES_NAME = "wheat_expansion_raw_pages.jsonl"
BATCH_META_NAME = "wheat_expansion_run_meta.json"
STEP0_ASIN_HEADER = "竞品ASIN"
SELLERSPRITE_ASIN_BATCH_SIZE = 20
SELLERSPRITE_EXPECTED_PAGE_EXTRA_LIMIT = 30
SELLERSPRITE_EMPTY_PAGE_AFTER_EXPECTED_LIMIT = 20
SELLERSPRITE_MARKETS = {
    "US": 1,
    "UK": 3,
    "DE": 4,
    "FR": 5,
    "IT": 35691,
    "ES": 44551,
    "CA": 7,
    "JP": 6,
    "MX": 771770,
}

@dataclass
class CurlRequest:
    url: str
    method: str
    headers: dict[str, str]
    body: dict[str, Any]


@dataclass(frozen=True)
class ConnectivityTestResult:
    sample_count: int
    total: int | None
    asin_batch_size: int


class WheatRequestInterrupted(RuntimeError):
    pass


def apply_wheat_market(body: dict[str, Any], country: str) -> dict[str, Any]:
    market = SELLERSPRITE_MARKETS.get(str(country or "").strip().upper())
    if market is None:
        raise RuntimeError("Step0 国家不受麦子支持。")
    updated = dict(body)
    updated["market"] = market
    return updated


def step0_wheat_country(path: str | Path, expected: str | None = None) -> str:
    from web_workbench.worker.steps.step0 import read_step0_market

    country = read_step0_market(Path(path))
    if expected is not None and str(expected).strip().upper() != country:
        raise RuntimeError("Step0 国家与当前任务市场不一致。")
    return country


def data_block(response: dict[str, Any]) -> dict[str, Any]:
    block = response.get("data")
    if not isinstance(block, dict):
        raise RuntimeError("JSON 里没有 data 对象。")
    return block


def safe_response_summary(response: dict[str, Any]) -> str:
    keys = ", ".join(sorted(str(key) for key in response.keys()))
    parts = [f"响应字段：{keys or '?'}"]
    code = response.get("code")
    if code is not None:
        parts.append(f"code={str(code)[:80]}")
    message = response.get("message")
    if message is not None:
        safe_message = str(message)[:200]
        for marker in ("cookie=", "authorization="):
            index = safe_message.lower().find(marker)
            if index >= 0:
                safe_message = safe_message[:index].rstrip()
        if safe_message:
            parts.append(f"message={safe_message}")
    success = response.get("success")
    if success is not None:
        parts.append(f"success={str(success)[:80]}")
    return "；".join(parts)


def data_block_for_page(response: dict[str, Any], context: str) -> dict[str, Any]:
    try:
        return data_block(response)
    except RuntimeError as exc:
        if "data 对象" not in str(exc):
            raise
        raise RuntimeError(f"{context}返回异常 JSON：缺少 data 对象；{safe_response_summary(response)}") from exc


def load_step0_competitor_asins(path: str | Path) -> list[str]:
    if openpyxl is None:
        raise RuntimeError("缺少 openpyxl，无法读取 Step0 竞品ASIN表。")
    workbook_path = Path(path)
    if not workbook_path.exists() or not workbook_path.is_file():
        raise RuntimeError(f"Step0 ASIN 文件不存在：{workbook_path}")
    if workbook_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise RuntimeError("Step0 ASIN 文件必须是 xlsx 或 xlsm。")

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = [str(value or "").strip() for value in (header_row or [])]
        if STEP0_ASIN_HEADER not in headers:
            raise RuntimeError(f"Step0 ASIN 文件缺少列：{STEP0_ASIN_HEADER}")
        asin_column = headers.index(STEP0_ASIN_HEADER)
        seen: set[str] = set()
        asins: list[str] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            value = row[asin_column] if asin_column < len(row) else ""
            asin = str(value or "").strip().upper()
            if not asin or asin in seen:
                continue
            seen.add(asin)
            asins.append(asin)
    finally:
        workbook.close()

    if not asins:
        raise RuntimeError("Step0 ASIN 文件中没有有效的竞品ASIN。")
    return asins


def batch_asins(asins: list[str], batch_size: int = SELLERSPRITE_ASIN_BATCH_SIZE) -> list[list[str]]:
    if batch_size <= 0:
        raise RuntimeError("ASIN 批次大小必须大于 0。")
    return [asins[index:index + batch_size] for index in range(0, len(asins), batch_size)]


def replace_wheat_asins(body: dict[str, Any], asins: list[str]) -> dict[str, Any]:
    if not asins:
        raise RuntimeError("ASIN 批次为空，无法请求麦子。")
    updated = copy.deepcopy(body)
    candidate_paths = [
        ("asins",),
        ("asinList",),
        ("asin_list",),
        ("query", "asins"),
        ("query", "asinList"),
        ("params", "asins"),
        ("params", "asinList"),
        ("data", "asins"),
        ("data", "asinList"),
    ]
    for path in candidate_paths:
        parent: Any = updated
        for key in path[:-1]:
            if not isinstance(parent, dict) or key not in parent:
                parent = None
                break
            parent = parent[key]
        if isinstance(parent, dict) and path[-1] in parent:
            parent[path[-1]] = list(asins)
            if path[-1] == "asinList" and "originAsinList" in parent:
                parent["originAsinList"] = list(asins)
            return updated
    raise RuntimeError("无法在麦子 cURL 请求体中找到 ASIN 列表字段，请重新复制 ASIN 拓展流量词接口的 cURL。")


def normalize_windows_curl(text: str) -> str:
    # Chrome on Windows copies cmd.exe-style cURL where ^ is an escape
    # character. Keep whitespace only for line continuations; remove the
    # remaining escape markers so ^\" becomes \" and ^& becomes &.
    return text.replace("^\r\n", " ").replace("^\n", " ").replace("^", "")


def parse_curl_file(path: str | Path) -> CurlRequest:
    raw = Path(path).read_text(encoding="utf-8-sig")
    normalized = normalize_windows_curl(raw)
    try:
        parts = shlex.split(normalized, posix=True)
    except ValueError as exc:
        raise RuntimeError(f"cURL 解析失败：{exc}") from exc
    if not parts or parts[0].lower() not in {"curl", "curl.exe"}:
        raise RuntimeError("文件内容不像 cURL：开头应为 curl。")

    url = ""
    method = "POST"
    headers: dict[str, str] = {}
    body_text = ""
    index = 1
    while index < len(parts):
        token = parts[index]
        if token in {"-H", "--header"}:
            index += 1
            if index >= len(parts):
                raise RuntimeError("cURL 里的 -H 缺少 header 内容。")
            header = parts[index]
            if ":" in header:
                name, value = header.split(":", 1)
                headers[name.strip()] = value.strip()
        elif token in {"-b", "--cookie", "--cookie-jar"}:
            index += 1
            if index >= len(parts):
                raise RuntimeError("cURL 里的 -b 缺少 cookie 内容。")
            headers["Cookie"] = parts[index]
        elif token in {"--data-raw", "--data", "--data-binary", "-d"}:
            index += 1
            if index >= len(parts):
                raise RuntimeError("cURL 里的 data 参数缺少内容。")
            body_text = parts[index]
            method = "POST"
        elif token in {"-X", "--request"}:
            index += 1
            if index >= len(parts):
                raise RuntimeError("cURL 里的 -X 缺少请求方法。")
            method = parts[index].upper()
        elif not token.startswith("-") and not url:
            url = token
        index += 1

    if not url:
        raise RuntimeError("cURL 里没有找到请求 URL。")
    if not body_text:
        raise RuntimeError("cURL 里没有找到 --data-raw，请确认复制的是后台 JSON 请求。")
    try:
        body = json.loads(body_text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"请求体 JSON 解析失败：第 {exc.lineno} 行第 {exc.colno} 列：{exc.msg}") from exc
    if not isinstance(body, dict):
        raise RuntimeError("请求体 JSON 顶层不是对象。")
    return CurlRequest(url=url, method=method, headers=headers, body=body)


def post_json(request: CurlRequest, body: dict[str, Any], timeout: int = 60) -> dict[str, Any]:
    data = json.dumps(body, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    headers = dict(request.headers)
    headers.setdefault("Content-Type", "application/json;charset=UTF-8")
    headers.setdefault("Accept", "application/json, text/plain, */*")
    req = urllib.request.Request(request.url, data=data, headers=headers, method=request.method)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as response:
            text = response.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        raise RuntimeError(f"HTTP {exc.code}: 麦子请求失败，请检查 cURL 是否过期或账号权限是否失效。") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"请求失败：{exc}") from exc
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"返回内容不是 JSON：第 {exc.lineno} 行第 {exc.colno} 列：{exc.msg}") from exc
    if not isinstance(parsed, dict):
        raise RuntimeError("接口返回 JSON 顶层不是对象。")
    return parsed


def load_existing_batch_pages(
    raw_pages_path: Path,
) -> tuple[set[int], dict[int, dict[str, Any]], int | None, int]:
    completed_batches: set[int] = set()
    batch_states: dict[int, dict[str, Any]] = {}
    page_size: int | None = None
    total_items = 0
    if not raw_pages_path.exists():
        return completed_batches, batch_states, page_size, total_items

    with raw_pages_path.open("r", encoding="utf-8") as raw_pages_file:
        for line_number, line in enumerate(raw_pages_file, start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise RuntimeError(f"已有批次原始数据第 {line_number} 行不是 JSON，无法断点续抓。") from exc
            if not isinstance(row, dict):
                continue
            batch = row.get("batch")
            page = row.get("page")
            request_data = row.get("request")
            if not isinstance(batch, int) or batch <= 0 or not isinstance(page, int) or page <= 0:
                raise RuntimeError(f"已有批次原始数据第 {line_number} 行缺少有效批次或页码，无法断点续抓。")
            if not isinstance(request_data, dict) or not isinstance(request_data.get("size"), int):
                raise RuntimeError(f"已有批次原始数据第 {line_number} 行缺少有效分页 size，无法断点续抓。")
            row_page_size = int(request_data["size"])
            if page_size is None:
                page_size = row_page_size
            elif row_page_size != page_size:
                raise RuntimeError("已有批次原始数据的分页 size 不一致，无法断点续抓。")

            state = batch_states.setdefault(batch, {
                "lastPage": 0,
                "rawReturnedItems": 0,
                "reportedTotal": None,
                "pages": 0,
                "consecutiveEmptyPages": 0,
            })
            if batch in completed_batches or page != int(state["lastPage"]) + 1:
                raise RuntimeError(f"已有批次原始数据第 {line_number} 行页码不连续，无法断点续抓。")
            response = row.get("response")
            if not isinstance(response, dict):
                raise RuntimeError(f"已有批次原始数据第 {line_number} 行缺少响应数据，无法断点续抓。")
            block = data_block(response)
            page_items = block.get("items") or []
            if not isinstance(page_items, list):
                raise RuntimeError(f"已有批次原始数据第 {line_number} 行 data.items 不是列表，无法断点续抓。")
            state["lastPage"] = page
            state["rawReturnedItems"] = int(state["rawReturnedItems"]) + len(page_items)
            total_items += len(page_items)
            state["pages"] = int(state["pages"]) + 1
            state["consecutiveEmptyPages"] = 0 if page_items else int(state["consecutiveEmptyPages"]) + 1
            if state["reportedTotal"] is None and isinstance(block.get("total"), int):
                state["reportedTotal"] = block.get("total")
            if row.get("batchComplete") is True:
                completed_batches.add(batch)
    return completed_batches, batch_states, page_size, total_items


def fetch_page_with_retries(
    request: CurlRequest,
    body: dict[str, Any],
    page: int,
    progress_callback: Callable[[str], None] | None = None,
    max_attempts: int = 3,
) -> dict[str, Any]:
    for attempt in range(1, max_attempts + 1):
        try:
            return post_json(request, body)
        except Exception as exc:
            if attempt >= max_attempts:
                raise WheatRequestInterrupted(f"第 {page} 页请求失败，已重试 {max_attempts} 次：{exc}") from exc
            wait_seconds = 5 * attempt
            message = f"第 {page} 页请求超时/失败，{wait_seconds} 秒后重试 {attempt}/{max_attempts - 1}：{exc}"
            print(message)
            if progress_callback:
                progress_callback(message)
            time.sleep(wait_seconds)
    raise RuntimeError(f"第 {page} 页请求失败。")


def fetch_wheat_asin_page(
    request: CurlRequest,
    body: dict[str, Any],
    page: int,
    progress_callback: Callable[[str], None] | None = None,
) -> dict[str, Any]:
    response = fetch_page_with_retries(request, body, page, progress_callback=progress_callback)
    block = response.get("data")
    if not isinstance(block, dict):
        return response
    expanded_asins = block.get("asinList")
    if (
        "items" not in block
        and "originAsinList" in body
        and isinstance(expanded_asins, list)
        and expanded_asins
        and all(isinstance(asin, str) and asin.strip() for asin in expanded_asins)
    ):
        prepared_body = copy.deepcopy(body)
        prepared_body["asinList"] = list(expanded_asins)
        return fetch_page_with_retries(request, prepared_body, page, progress_callback=progress_callback)
    return response


def test_curl_and_asins_connectivity(
    curl_path: str | Path,
    asin_workbook_path: str | Path,
    progress_callback: Callable[[str], None] | None = None,
    *,
    country: str | None = None,
) -> ConnectivityTestResult:
    request = parse_curl_file(curl_path)
    country = step0_wheat_country(asin_workbook_path, country)
    body_template = apply_wheat_market(request.body, country)
    page_size = int(body_template.get("size") or 100)
    if page_size <= 0:
        raise RuntimeError("分页 size 必须大于 0。")

    asins = load_step0_competitor_asins(asin_workbook_path)
    batches = batch_asins(asins)
    if not batches:
        raise RuntimeError("Step0 竞品ASIN表没有可用 ASIN，无法连通测试。")

    asin_batch = batches[0]
    body = replace_wheat_asins(body_template, asin_batch)
    body["page"] = 1
    body["size"] = page_size
    response = fetch_wheat_asin_page(request, body, 1, progress_callback=progress_callback)
    block = data_block(response)
    page_items = block.get("items") or []
    if not isinstance(page_items, list):
        raise RuntimeError("连通测试失败：接口返回 data.items 不是列表。")
    if not page_items:
        raise RuntimeError("连通测试失败：接口返回 data.items 为空，请检查 cURL、账号权限或 Step0 ASIN。")

    has_keyword_data = any(isinstance(item, dict) and str(item.get("keywords") or "").strip() for item in page_items)
    if not has_keyword_data:
        raise RuntimeError("连通测试失败：样例数据中没有 keywords 字段，请确认复制的是麦子拓展流量词接口 cURL。")

    total = block.get("total") if isinstance(block.get("total"), int) else None
    message = f"连通测试第1批第1页：返回 {len(page_items)} 条样例数据，总量约 {total or '?'}。"
    print(message)
    if progress_callback:
        progress_callback(message)
    return ConnectivityTestResult(sample_count=len(page_items), total=total, asin_batch_size=len(asin_batch))


def fetch_from_curl_and_asins(
    curl_path: str | Path,
    asin_workbook_path: str | Path,
    output_dir: str | Path | None = None,
    max_batches: int | None = None,
    delay: float = 1.2,
    progress_callback: Callable[[str], None] | None = None,
    resume: bool = False,
    resume_checkpoint: dict[str, object] | None = None,
    save_checkpoint: Callable[[dict[str, object]], None] | None = None,
    stop_requested: Callable[[], bool] | None = None,
    publish_progress: Callable[[int, int], None] | None = None,
    country: str | None = None,
) -> tuple[Path, Path, int]:
    request = parse_curl_file(curl_path)
    country = step0_wheat_country(asin_workbook_path, country)
    body_template = apply_wheat_market(request.body, country)
    page_size = int(body_template.get("size") or 100)
    if page_size <= 0:
        raise RuntimeError("分页 size 必须大于 0。")

    asins = load_step0_competitor_asins(asin_workbook_path)
    batches = batch_asins(asins)
    target_batches = batches[:max_batches] if max_batches is not None else batches
    total_batches = len(target_batches)
    request_fingerprint = credential_free_curl_fingerprint(Path(curl_path).read_text(encoding="utf-8"))
    asin_fingerprint = hashlib.sha256(
        json.dumps(asins, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    ).hexdigest()

    target_dir = Path(output_dir) if output_dir else create_run_output_dir()
    target_dir.mkdir(parents=True, exist_ok=True)
    raw_pages_path = target_dir / RAW_PAGES_NAME
    meta_path = target_dir / BATCH_META_NAME

    completed_batches: set[int] = set()
    batch_states: dict[int, dict[str, Any]] = {}
    total_raw_items = 0
    file_mode = "w"
    if resume:
        completed_batches, batch_states, existing_page_size, total_raw_items = load_existing_batch_pages(raw_pages_path)
        if existing_page_size is not None and existing_page_size != page_size:
            raise RuntimeError("当前分页 size 与断点数据不一致，无法断点续抓。")
        expected_checkpoint = {
            "pageSize": page_size,
            "totalBatches": total_batches,
            "requestFingerprint": request_fingerprint,
            "asinFingerprint": asin_fingerprint,
        }
        for key, expected in expected_checkpoint.items():
            actual = (resume_checkpoint or {}).get(key)
            if actual is not None and actual != expected:
                raise RuntimeError("当前麦子请求或 Step0 输入与断点不一致，无法断点续抓。")
        file_mode = "a" if raw_pages_path.exists() else "w"
        if completed_batches:
            message = f"断点续抓：已完成批次 {sorted(completed_batches)}，已有 {total_raw_items} 条原始记录。"
            print(message)
            if progress_callback:
                progress_callback(message)

    def checkpoint_page(current_batch: int, current_page: int, last_saved_batch: int, last_saved_page: int) -> None:
        if save_checkpoint:
            save_checkpoint({
                "completedBatch": max(completed_batches, default=0),
                "currentBatch": current_batch,
                "currentPage": current_page,
                "lastSavedBatch": last_saved_batch,
                "lastSavedPage": last_saved_page,
                "totalBatches": total_batches,
                "pageSize": page_size,
                "outputDir": str(target_dir),
                "requestFingerprint": request_fingerprint,
                "asinFingerprint": asin_fingerprint,
            })

    if resume and batch_states:
        last_saved_batch = max(batch_states)
        last_saved_page = int(batch_states[last_saved_batch]["lastPage"])
        if last_saved_batch in completed_batches:
            current_batch, current_page = min(last_saved_batch + 1, total_batches), 1
        else:
            current_batch, current_page = last_saved_batch, last_saved_page + 1
        checkpoint_page(current_batch, current_page, last_saved_batch, last_saved_page)
    elif not completed_batches:
        checkpoint_page(1, 1, 0, 0)

    batch_stats: list[dict[str, Any]] = []
    stopped_mid_batch = False
    with raw_pages_path.open(file_mode, encoding="utf-8") as raw_pages_file:
        for batch_index, asin_batch in enumerate(target_batches, start=1):
            if batch_index in completed_batches:
                continue
            if stop_requested and stop_requested():
                break
            existing_state = batch_states.get(batch_index, {})
            page = int(existing_state.get("lastPage") or 0) + 1
            raw_returned_items = int(existing_state.get("rawReturnedItems") or 0)
            batch_total = existing_state.get("reportedTotal")
            batch_total = batch_total if isinstance(batch_total, int) else None
            expected_pages: int | None = None
            max_pages: int | None = None
            if batch_total is not None and batch_total > 0:
                expected_pages = max(1, math.ceil(batch_total / page_size))
                max_pages = expected_pages + SELLERSPRITE_EXPECTED_PAGE_EXTRA_LIMIT
            batch_pages = int(existing_state.get("pages") or 0)
            consecutive_empty_pages = int(existing_state.get("consecutiveEmptyPages") or 0)
            stop_reason = ""
            warning = ""
            while True:
                body = replace_wheat_asins(body_template, asin_batch)
                body["page"] = page
                body["size"] = page_size
                response = fetch_wheat_asin_page(request, body, page, progress_callback=progress_callback)
                block = data_block_for_page(response, f"第 {batch_index} 批第 {page} 页")
                page_items = block.get("items") or []
                if not isinstance(page_items, list):
                    raise RuntimeError(f"第 {batch_index} 批第 {page} 页返回 data.items 不是列表。")
                batch_pages += 1
                if page_items:
                    consecutive_empty_pages = 0
                else:
                    consecutive_empty_pages += 1
                if batch_total is None and isinstance(block.get("total"), int):
                    batch_total = block.get("total")
                    if batch_total > 0:
                        expected_pages = max(1, math.ceil(batch_total / page_size))
                        max_pages = expected_pages + SELLERSPRITE_EXPECTED_PAGE_EXTRA_LIMIT

                raw_returned_items += len(page_items)
                total_raw_items += len(page_items)

                if batch_total is not None:
                    if raw_returned_items >= batch_total:
                        is_complete = True
                        stop_reason = "reached_total"
                    elif (
                        expected_pages is not None
                        and page >= expected_pages
                        and consecutive_empty_pages >= SELLERSPRITE_EMPTY_PAGE_AFTER_EXPECTED_LIMIT
                    ):
                        is_complete = True
                        stop_reason = "empty_pages_after_expected_pages"
                        warning = "reported total not reached"
                    elif max_pages is not None and page >= max_pages:
                        is_complete = True
                        stop_reason = "exceeded_expected_pages"
                        warning = "reported total not reached"
                    else:
                        is_complete = False
                else:
                    if not page_items:
                        is_complete = True
                        stop_reason = "empty_page_no_total"
                    elif len(page_items) < page_size:
                        is_complete = True
                        stop_reason = "short_page_no_total"
                    else:
                        is_complete = False
                raw_pages_file.write(json.dumps({
                    "batch": batch_index,
                    "page": page,
                    "request": {"batch": batch_index, "page": page, "size": page_size, "asins": asin_batch},
                    "response": response,
                    "batchComplete": is_complete,
                }, ensure_ascii=False))
                raw_pages_file.write("\n")
                raw_pages_file.flush()
                os.fsync(raw_pages_file.fileno())

                if is_complete:
                    completed_batches.add(batch_index)
                    next_batch = min(batch_index + 1, total_batches)
                    checkpoint_page(next_batch, 1, batch_index, page)
                else:
                    checkpoint_page(batch_index, page + 1, batch_index, page)

                message = (
                    f"第 {batch_index}/{len(target_batches)} 批第 {page} 页：返回 {len(page_items)} 条，"
                    f"本批原始累计 {raw_returned_items}/{batch_total or '?'}"
                )
                print(message)
                if progress_callback:
                    progress_callback(message)

                if not is_complete and stop_requested and stop_requested():
                    stopped_mid_batch = True
                    break

                if is_complete:
                    batch_stat: dict[str, Any] = {
                        "batch": batch_index,
                        "reportedTotal": batch_total,
                        "rawReturnedItems": raw_returned_items,
                        "pages": batch_pages,
                        "expectedPages": expected_pages,
                        "maxPages": max_pages,
                        "stopReason": stop_reason,
                    }
                    if warning:
                        batch_stat["warning"] = warning
                        warn_message = (
                            f"警告：第 {batch_index} 批接口 reported total={batch_total}，"
                            f"原始返回 {raw_returned_items} 条，"
                            f"stopReason={stop_reason}。"
                        )
                        print(warn_message)
                        if progress_callback:
                            progress_callback(warn_message)
                    batch_stats.append(batch_stat)
                    if publish_progress:
                        publish_progress(batch_index, total_batches)
                    break
                page += 1
                if delay > 0:
                    time.sleep(delay)
            if stopped_mid_batch:
                break

    warnings = [
        f"第 {stat['batch']} 批 reported total not reached: rawReturnedItems={stat['rawReturnedItems']}/{stat['reportedTotal']}"
        for stat in batch_stats
        if stat.get("warning")
    ]
    meta = {
        "sourceCurlFile": str(curl_path),
        "sourceAsinWorkbook": str(asin_workbook_path),
        "asinCount": len(asins),
        "batchSize": SELLERSPRITE_ASIN_BATCH_SIZE,
        "batchCount": len(batches),
        "completedBatches": sorted(completed_batches),
        "batchStats": batch_stats,
        "warnings": warnings,
    }
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return raw_pages_path, meta_path, total_raw_items


def main() -> None:
    parser = argparse.ArgumentParser(description="麦子 ASIN 拓展原始 JSONL 下载工具")
    parser.add_argument("curl_file", help="保存 Copy as cURL 内容的 txt 文件")
    parser.add_argument("asin_workbook", help="ASIN 输入表 xlsx")
    parser.add_argument("--output-dir", help="输出目录；默认写到 output/北京时间戳/")
    parser.add_argument("--max-batches", type=int, help="最多抓多少个 ASIN 批次")
    parser.add_argument("--delay", type=float, default=1.2, help="每页间隔秒数")
    parser.add_argument("--resume", action="store_true", help="从已有 JSONL 继续下载")
    args = parser.parse_args()
    raw_pages, meta_path, rows = fetch_from_curl_and_asins(
        args.curl_file,
        args.asin_workbook,
        args.output_dir,
        args.max_batches,
        args.delay,
        resume=args.resume,
    )
    print(f"下载完成：{rows} 条原始记录")
    print(f"分页原始数据：{raw_pages}")
    print(f"运行元信息：{meta_path}")


if __name__ == "__main__":
    main()
