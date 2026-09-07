"""Output path, backup, manifest, and xlsx validation helpers."""

from __future__ import annotations

import csv
import json
import os
import re
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone, timedelta
from pathlib import Path
from uuid import uuid4

import product_paths

try:
    import openpyxl
except Exception:  # pragma: no cover - dependency is checked at runtime
    openpyxl = None

ROOT = product_paths.app_root()
CONFIG_PATH = product_paths.resource_path("config.json")
BEIJING_TZ = timezone(timedelta(hours=8))

LEGACY_MANIFEST_FIELDS = [
    "run_time",
    "country",
    "asin",
    "status",
    "output_file",
    "error",
    "attempts",
]
MANIFEST_FIELDS = [*LEGACY_MANIFEST_FIELDS, "download_diagnostics"]


STEP_DIRS = {
    "step0_competitor_asins": "Step0_ASIN_Input",
    "Step1_Pomelo_Data": "Step1_Pomelo_Data",
    "Step2_Wheat_Expansion": "Step2_Wheat_Expansion",
    "Step9_Pomelo_Keywords": "Step9_Pomelo_Keywords",
}

STEP0_COMPETITOR_WORKBOOK_NAME = "Step0_ASIN_Input.xlsx"
STEP0_COMPETITOR_HEADERS = ["国家", "竞品ASIN", "竞品强弱"]
EXCEL_ILLEGAL_CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F\uD800-\uDFFF\uFFFE\uFFFF]")


def sanitize_excel_value(value: object) -> object:
    return EXCEL_ILLEGAL_CONTROL_CHARACTERS.sub("", value) if isinstance(value, str) else value


def append_excel_row(sheet: object, row: object) -> None:
    sheet.append([sanitize_excel_value(value) for value in row])


def step0_competitor_workbook_path(task_dir: str | Path) -> Path:
    return Path(task_dir) / STEP_DIRS["step0_competitor_asins"] / STEP0_COMPETITOR_WORKBOOK_NAME


def ensure_step0_competitor_template(task_dir: str | Path) -> Path:
    template_path = step0_competitor_workbook_path(task_dir)
    if template_path.exists():
        return template_path
    if openpyxl is None:
        raise RuntimeError("缺少 openpyxl，无法创建 Step0 竞品ASIN模板。")
    workbook = openpyxl.Workbook()
    try:
        sheet = workbook.active
        sheet.title = "Step0竞品ASIN"
        sheet.append(STEP0_COMPETITOR_HEADERS)
        workbook.save(template_path)
    finally:
        workbook.close()
    return template_path


def sanitize_task_slug(value: str) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text or "task"


def create_task_dir(task_name: str, config: dict | None = None, timestamp: str | None = None) -> Path:
    root = output_root(config)
    date_prefix = (timestamp or beijing_timestamp())[:8]
    slug = sanitize_task_slug(task_name)
    base_name = f"task_{date_prefix}_{slug}"
    candidate = root / base_name
    if not candidate.exists():
        ensure_task_structure(candidate)
        return candidate
    for index in range(2, 100):
        candidate = root / f"{base_name}_{index:02d}"
        if not candidate.exists():
            ensure_task_structure(candidate)
            return candidate
    raise RuntimeError(f"无法创建任务目录：{root / base_name}")


def ensure_task_structure(task_dir: str | Path) -> Path:
    task_path = Path(task_dir)
    task_path.mkdir(parents=True, exist_ok=True)
    for dirname in STEP_DIRS.values():
        (task_path / dirname).mkdir(parents=True, exist_ok=True)
    ensure_step0_competitor_template(task_path)
    request_files = {
        "Step2_Wheat_Expansion": "Step2_Request.txt",
        "Step9_Pomelo_Keywords": "Step9_Request.txt",
    }
    for step_key, filename in request_files.items():
        request_path = task_path / STEP_DIRS[step_key] / filename
        if not request_path.exists():
            request_path.write_text("", encoding="utf-8")
    return task_path


def step_dir(task_dir: str | Path, step_key: str) -> Path:
    if step_key not in STEP_DIRS:
        raise RuntimeError(f"未知步骤目录：{step_key}")
    return ensure_task_structure(task_dir) / STEP_DIRS[step_key]


def create_step_run_dir(task_dir: str | Path, step_key: str, prefix: str = "", timestamp: str | None = None) -> Path:
    parent = step_dir(task_dir, step_key)
    base_name = f"{prefix}{timestamp or beijing_timestamp()}"
    candidate = parent / base_name
    if not candidate.exists():
        candidate.mkdir(parents=True)
        return candidate
    for index in range(2, 100):
        candidate = parent / f"{base_name}_{index:02d}"
        if not candidate.exists():
            candidate.mkdir(parents=True)
            return candidate
    raise RuntimeError(f"无法创建步骤输出目录：{parent / base_name}")


@dataclass(frozen=True)
class ValidationResult:
    ok: bool
    rows: int = 0
    columns: int = 0
    error: str = ""


def load_config() -> dict:
    with CONFIG_PATH.open("r", encoding="utf-8") as file:
        return json.load(file)


def output_root(config: dict | None = None) -> Path:
    cfg = config or load_config()
    path = ROOT / cfg.get("output_dir", "output")
    path.mkdir(parents=True, exist_ok=True)
    return path


def output_dir(config: dict | None = None) -> Path:
    """Return the base output folder; kept for UI/open-folder compatibility."""
    return output_root(config)


def beijing_timestamp() -> str:
    return datetime.now(BEIJING_TZ).strftime("%Y%m%d_%H%M%S")


def create_run_output_dir(config: dict | None = None, timestamp: str | None = None) -> Path:
    root = output_root(config)
    base_name = timestamp or beijing_timestamp()
    candidate = root / base_name
    if not candidate.exists():
        candidate.mkdir(parents=True)
        return candidate
    for index in range(2, 100):
        candidate = root / f"{base_name}_{index:02d}"
        if not candidate.exists():
            candidate.mkdir(parents=True)
            return candidate
    raise RuntimeError(f"无法创建输出目录：{root / base_name}")


def ensure_run_output_dir(run_output_dir: str | Path | None = None, config: dict | None = None) -> Path:
    if run_output_dir is None:
        return create_run_output_dir(config)
    path = Path(run_output_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def keyword_output_name(country: str, asin: str) -> str:
    safe_country = country.strip().upper()
    safe_asin = asin.strip().upper()
    return f"{safe_country}_{safe_asin}_关键词反查.xlsx"


def keyword_output_path(country: str, asin: str, config: dict | None = None, run_output_dir: str | Path | None = None) -> Path:
    return ensure_run_output_dir(run_output_dir, config) / keyword_output_name(country, asin)


def manifest_path(config: dict | None = None, run_output_dir: str | Path | None = None) -> Path:
    return ensure_run_output_dir(run_output_dir, config) / "批量结果清单.csv"


def backup_existing_file(path: str | Path) -> Path | None:
    target = Path(path)
    if not target.exists():
        return None
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = target.with_name(f"{target.stem}.old-{timestamp}{target.suffix}")
    shutil.move(str(target), str(backup))
    return backup


def validate_keyword_xlsx(path: str | Path) -> ValidationResult:
    file_path = Path(path)
    if not file_path.exists():
        return ValidationResult(False, error="文件不存在")
    if file_path.stat().st_size <= 0:
        return ValidationResult(False, error="下载文件为空")
    if openpyxl is None:
        return ValidationResult(False, error="缺少 openpyxl，无法校验 xlsx")
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.max_row or 0
            columns = sheet.max_column or 0
            if rows < 2:
                return ValidationResult(False, rows, columns, "xlsx 只有表头无数据")
            if columns < 1:
                return ValidationResult(False, rows, columns, "xlsx 没有有效列")
            first_data = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), None)
            if not first_data or not any(str(cell or "").strip() for cell in first_data):
                return ValidationResult(False, rows, columns, "xlsx 第一行数据为空")
            return ValidationResult(True, rows, columns, "")
        finally:
            workbook.close()
    except Exception as exc:
        return ValidationResult(False, error=f"xlsx 无法打开：{exc}")


def _read_manifest_rows(path: Path) -> tuple[list[str], list[dict[str, str | None]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        fields = reader.fieldnames
        if fields not in (LEGACY_MANIFEST_FIELDS, MANIFEST_FIELDS):
            raise ValueError("pomelo_manifest_schema_invalid")
        rows = list(reader)
    if any(None in row for row in rows):
        raise ValueError("pomelo_manifest_schema_invalid")
    return fields, rows


def _upgrade_legacy_manifest(path: Path, rows: list[dict[str, str | None]]) -> None:
    staged = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    try:
        with staged.open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=MANIFEST_FIELDS, extrasaction="raise")
            writer.writeheader()
            for row in rows:
                writer.writerow({**row, "download_diagnostics": "[]"})
        staged_fields, staged_rows = _read_manifest_rows(staged)
        if staged_fields != MANIFEST_FIELDS or len(staged_rows) != len(rows):
            raise ValueError("pomelo_manifest_schema_invalid")
        os.replace(staged, path)
    finally:
        staged.unlink(missing_ok=True)


def append_manifest_row(row: dict, config: dict | None = None, run_output_dir: str | Path | None = None) -> Path:
    path = manifest_path(config, run_output_dir)
    exists = path.exists()
    if exists:
        fields, rows = _read_manifest_rows(path)
        if fields == LEGACY_MANIFEST_FIELDS:
            _upgrade_legacy_manifest(path, rows)
    with path.open("a", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=MANIFEST_FIELDS, extrasaction="raise")
        if not exists:
            writer.writeheader()
        writer.writerow(row)
    return path
