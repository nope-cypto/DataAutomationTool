from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from input_parser import SUPPORTED_COUNTRIES, normalize_country
from web_workbench.worker.files import is_transient_workflow_file


REQUIRED_COLUMNS = ["国家", "竞品ASIN", "竞品强弱"]


def validate_step0_workbook(path: Path) -> dict:
    if not path.exists() or path.suffix.lower() != ".xlsx" or is_transient_workflow_file(path):
        return {"ok": False, "error": "step0_workbook_missing", "missingColumns": REQUIRED_COLUMNS, "asinCount": 0}

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1))
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header_row]
        missing = [column for column in REQUIRED_COLUMNS if column not in headers]
        if missing:
            return {
                "ok": False,
                "error": "step0_columns_missing",
                "missingColumns": missing,
                "asinCount": 0,
            }

        asin_count = 0
        asin_index = headers.index("竞品ASIN")
        country_index = headers.index("国家")
        countries: set[str] = set()
        missing_country = False
        unsupported_country = False
        for row in sheet.iter_rows(min_row=2, values_only=True):
            asin = str(row[asin_index] if asin_index < len(row) else "").strip()
            if not asin:
                continue
            asin_count += 1
            country = normalize_country(row[country_index] if country_index < len(row) else "")
            if not country:
                missing_country = True
            elif country not in SUPPORTED_COUNTRIES:
                unsupported_country = True
            else:
                countries.add(country)

        error = None
        if asin_count == 0:
            error = "step0_asin_missing"
        elif missing_country:
            error = "step0_country_missing"
        elif unsupported_country:
            error = "step0_country_unsupported"
        elif len(countries) != 1:
            error = "step0_country_mixed"
        if error:
            return {
                "ok": False,
                "error": error,
                "missingColumns": [],
                "asinCount": asin_count,
            }
        return {
            "ok": True,
            "missingColumns": [],
            "asinCount": asin_count,
            "country": next(iter(countries)),
        }
    finally:
        workbook.close()


def read_step0_market(path: Path) -> str:
    result = validate_step0_workbook(path)
    if not result.get("ok"):
        raise ValueError(str(result.get("error") or "step0_market_invalid"))
    return str(result["country"])
